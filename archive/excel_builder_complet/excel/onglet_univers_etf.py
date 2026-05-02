import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.excel.styles import (
    COULEURS_CLASSES,
    _align,
    _fill,
    _font,
    _thin_border,
    ajouter_disclaimer,
    set_col_width,
    style_header,
    style_subheader,
)


# ─── Onglet 4 : Univers ETF ─────────────────────────────────────────
def creer_onglet_univers_etf(wb: openpyxl.Workbook, etfs: list):
    ws = wb.create_sheet("Univers_ETF")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:N1")
    c = ws["A1"]
    c.value = f"📈 UNIVERS ETF — BOGLEHEAD FR 2026 ({len(etfs)} ETF)"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30

    ajouter_disclaimer(ws, 2, 1, 14)

    headers = [
        "ISIN",
        "Ticker",
        "Nom",
        "Émetteur",
        "Classe d'actifs",
        "Sous-classe",
        "TER (%)",
        "Devise",
        "Capitalisant",
        "EUR-Hedgé",
        "PEA",
        "PER",
        "CTO",
        "Notes",
    ]
    for i, h in enumerate(headers, 1):
        style_subheader(ws.cell(row=3, column=i, value=h))
    ws.row_dimensions[3].height = 22

    for idx, etf in enumerate(etfs):
        row = idx + 4
        classe = etf.get("classe_actifs", "")
        bg_hex = COULEURS_CLASSES.get(classe, "FFFFFF")
        # Alternance légère
        bg = bg_hex if idx % 2 == 0 else "FFFFFF"

        elig = etf.get("eligibilite", {})
        row_data = [
            etf.get("isin", ""),
            etf.get("ticker", ""),
            etf.get("nom", ""),
            etf.get("emetteur", ""),
            classe,
            etf.get("sous_classe", ""),
            etf.get("ter", 0),
            etf.get("devise", ""),
            "✅" if etf.get("capitalisant") else "❌",
            "✅" if etf.get("eur_hedged") else "❌",
            "✅" if elig.get("PEA") else "❌",
            "✅" if elig.get("PER") else "❌",
            "✅" if elig.get("CTO_perso") else "❌",
            etf.get("notes", ""),
        ]
        for j, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(size=9)
            cell.alignment = _align(
                "center" if j in (7, 8, 9, 10, 11, 12, 13) else "left", wrap=True
            )
            cell.border = _thin_border()
            if j == 7:  # TER
                cell.number_format = "0.00%"

    # Tableau Excel structuré
    n_rows = len(etfs)
    if n_rows > 0:
        last_row = n_rows + 3
        tbl = Table(displayName="tblETF", ref=f"A3:N{last_row}")
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tbl.tableStyleInfo = style
        ws.add_table(tbl)

    widths = [16, 8, 50, 20, 15, 25, 8, 8, 10, 10, 6, 6, 6, 45]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:N{len(etfs) + 3}"
