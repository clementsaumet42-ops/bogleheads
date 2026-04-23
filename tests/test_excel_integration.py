"""Integration test: generate Excel file and verify structure."""

import os
from pathlib import Path

import openpyxl
import pytest

OUTPUT_PATH = Path(__file__).parent.parent / "output" / "test_bogleheads.xlsx"


@pytest.fixture(scope="module")
def workbook():
    """Generate Excel file once for all tests in module."""
    from src.excel_builder import generer_excel

    OUTPUT_PATH.parent.mkdir(exist_ok=True)
    generer_excel(str(OUTPUT_PATH))
    assert OUTPUT_PATH.exists(), "Excel file was not created"
    wb = openpyxl.load_workbook(str(OUTPUT_PATH))
    yield wb
    wb.close()


class TestExcelStructure:
    EXPECTED_SHEETS = [
        "Paramètres_Client",
        "Paramètres_Fiscalité_2026",
        "Enveloppes",
        "Univers_ETF",
        "Allocation_Cible",
        "Asset_Location_Matrice",
        "Rebalancement",
        "Reporting_Client",
    ]

    def test_fichier_cree(self):
        assert OUTPUT_PATH.exists()
        assert OUTPUT_PATH.stat().st_size > 0

    def test_onglets_presents(self, workbook):
        sheet_names = workbook.sheetnames
        for expected in self.EXPECTED_SHEETS:
            assert expected in sheet_names, f"Missing sheet: '{expected}'"

    def test_onglet_parametres_non_vide(self, workbook):
        ws = workbook["Paramètres_Client"]
        non_empty = sum(1 for row in ws.iter_rows() for cell in row if cell.value)
        assert non_empty > 5

    def test_onglet_fiscalite_non_vide(self, workbook):
        ws = workbook["Paramètres_Fiscalité_2026"]
        non_empty = sum(1 for row in ws.iter_rows() for cell in row if cell.value)
        assert non_empty > 5

    def test_onglet_enveloppes_contient_pea(self, workbook):
        ws = workbook["Enveloppes"]
        texts = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        assert any("PEA" in t for t in texts)

    def test_onglet_univers_etf_contient_isin(self, workbook):
        ws = workbook["Univers_ETF"]
        texts = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        assert any("IE0031442068" in t or "CW8" in t for t in texts)

    def test_onglet_allocation_presente(self, workbook):
        ws = workbook["Allocation_Cible"]
        non_empty = sum(1 for row in ws.iter_rows() for cell in row if cell.value)
        assert non_empty > 3

    def test_20_onglets_crees(self, workbook):
        assert len(workbook.sheetnames) >= 20
