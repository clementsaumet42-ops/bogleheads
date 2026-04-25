from __future__ import annotations

from src.audit.alertes.base import Severite

_COULEUR_ROUGE = "FFCCCC"
_COULEUR_JAUNE = "FFF9CC"
_COULEUR_VERT = "CCFFCC"

_COULEUR_MAP = {
    Severite.ROUGE: _COULEUR_ROUGE,
    Severite.JAUNE: _COULEUR_JAUNE,
    Severite.VERT: _COULEUR_VERT,
}


def creer_onglet_alertes(ws, alertes, toutes_les_regles=None):
    """Crée l'onglet 'Alertes (40 règles)' dans le classeur Excel."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        Font = PatternFill = Alignment = None

    headers = [
        "Code", "Famille", "Sévérité", "Titre", "Description",
        "Gain €/an", "Gain € horizon", "Action concrète", "Sources", "Statut",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        if Font and PatternFill:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="1B3A5B")
            cell.font = Font(bold=True, color="FFFFFF")

    for row, a in enumerate(alertes, 2):
        ws.cell(row=row, column=1, value=a.code)
        ws.cell(row=row, column=2, value=a.famille)
        ws.cell(row=row, column=3, value=str(a.severite.value))
        ws.cell(row=row, column=4, value=a.titre)
        ws.cell(row=row, column=5, value=a.description)
        ws.cell(row=row, column=6, value=a.gain_eur_annuel)
        ws.cell(row=row, column=7, value=a.gain_eur_horizon)
        ws.cell(row=row, column=8, value=a.action_concrete)
        ws.cell(row=row, column=9, value="; ".join(a.sources))
        ws.cell(row=row, column=10, value="Déclenchée")

        if PatternFill:
            couleur = _COULEUR_MAP.get(a.severite, "FFFFFF")
            fill = PatternFill("solid", fgColor=couleur)
            for col in range(1, 11):
                ws.cell(row=row, column=col).fill = fill

    # Auto-width approximation
    col_widths = [8, 22, 10, 45, 60, 12, 14, 60, 40, 12]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
