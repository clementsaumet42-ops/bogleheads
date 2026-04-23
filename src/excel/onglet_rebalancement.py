import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from src.excel.styles import (
    COULEUR_HEADER, COULEUR_SUBHEADER, COULEUR_OK, COULEUR_WARNING, COULEUR_DANGER, COULEUR_LIGHT_GREY,
    _fill, _font, _align, _thin_border,
    style_subheader, set_col_width, ajouter_disclaimer, titre_section, style_header,
)


# ─── Onglet 7 : Rebalancement ───────────────────────────────────────
def creer_onglet_rebalancement(wb: openpyxl.Workbook):
    ws = wb.create_sheet("Rebalancement")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "⚖️ OUTIL DE REBALANCEMENT — BOGLEHEAD FR"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30
    ajouter_disclaimer(ws, 2, 1, 8)

    row = 4
    row = titre_section(ws, row, "📊 TABLEAU DE SUIVI DE L'ALLOCATION", 1, 8)
    headers = ["Classe d'actifs", "Cible (%)", "Actuelle (%)", "Dérive (pts abs.)",
               "Dérive (%rel)", "Action", "Montant à arbitrer (€)", "Enveloppe privilégiée"]
    for i, h in enumerate(headers, 1):
        style_subheader(ws.cell(row=row, column=i, value=h))
    row += 1

    classes = ["Actions", "Obligations", "Immobilier coté", "Or", "Liquidités", "TOTAL"]
    cibles = [0.65, 0.20, 0.05, 0.05, 0.05, 1.00]
    env_pref = ["PEA/CTO", "PER/ContratIS", "PER/CTO", "CTO", "CTO", ""]
    patrimoine_exemple = 500000

    for i, (classe, cible, env) in enumerate(zip(classes, cibles, env_pref)):
        is_total = classe == "TOTAL"
        bg = COULEUR_HEADER if is_total else (COULEUR_LIGHT_GREY if i % 2 == 0 else "FFFFFF")
        fg = "FFFFFF" if is_total else "000000"

        ws.cell(row=row, column=1, value=classe).fill = _fill(bg)
        ws.cell(row=row, column=1).font = _font(bold=is_total, color=fg)
        ws.cell(row=row, column=2, value=cible).number_format = "0.0%"
        ws.cell(row=row, column=2).fill = _fill(bg)

        if not is_total:
            # Formule dérive = actuelle - cible
            derive_col = get_column_letter(4)
            actuelle_col = get_column_letter(3)
            cible_col = get_column_letter(2)

            ws.cell(row=row, column=3, value=cible).number_format = "0.0%"  # à saisir
            ws.cell(row=row, column=3).fill = _fill("FFFFC0")  # Zone de saisie
            ws.cell(row=row, column=3).comment = None

            ws.cell(row=row, column=4,
                    value=f"={actuelle_col}{row}-{cible_col}{row}").number_format = "0.0%"
            ws.cell(row=row, column=4).fill = _fill(bg)

            ws.cell(row=row, column=5,
                    value=f"=IF({cible_col}{row}>0,ABS({derive_col}{row})/{cible_col}{row},0)").number_format = "0.0%"
            ws.cell(row=row, column=5).fill = _fill(bg)

            ws.cell(row=row, column=6,
                    value=f'=IF(ABS({derive_col}{row})>0.05,"ARBITRER",IF(E{row}>0.25,"SURVEILLER","OK"))')
            ws.cell(row=row, column=6).alignment = _align("center")

            ws.cell(row=row, column=7,
                    value=f"=ABS({derive_col}{row})*{patrimoine_exemple}").number_format = "#,##0 €"
            ws.cell(row=row, column=7).fill = _fill(bg)

            ws.cell(row=row, column=8, value=env).font = _font(italic=True)
        else:
            ws.cell(row=row, column=3, value=f"=SUM(C{row-5}:C{row-1})").number_format = "0.0%"
            for col in range(2, 9):
                ws.cell(row=row, column=col).fill = _fill(bg)
                ws.cell(row=row, column=col).font = Font(bold=True, color=fg)

        for col in range(1, 9):
            ws.cell(row=row, column=col).border = _thin_border()
        row += 1

    # Mise en forme conditionnelle sur la colonne Action (F)
    green_fill = PatternFill(bgColor=COULEUR_OK, fill_type="solid")
    red_fill = PatternFill(bgColor=COULEUR_DANGER, fill_type="solid")
    orange_fill = PatternFill(bgColor=COULEUR_WARNING, fill_type="solid")
    action_range = f"F{row-5}:F{row-1}"
    ws.conditional_formatting.add(
        action_range,
        CellIsRule(operator="equal", formula=['"OK"'], fill=green_fill)
    )
    ws.conditional_formatting.add(
        action_range,
        CellIsRule(operator="equal", formula=['"ARBITRER"'], fill=red_fill)
    )
    ws.conditional_formatting.add(
        action_range,
        CellIsRule(operator="equal", formula=['"SURVEILLER"'], fill=orange_fill)
    )

    row += 1
    row = titre_section(ws, row, "📋 MÉTHODE DE REBALANCEMENT — ORDRE DE PRIORITÉ", 1, 8)
    methodes = [
        ("1ère priorité : Versements", "Orienter les nouveaux versements vers les classes sous-pondérées → coût fiscal = 0"),
        ("2ème priorité : Arbitrage intra-enveloppe", "PEA, PER, PEE → arbitrage SANS fiscalité (enveloppe = bouclier fiscal)"),
        ("3ème priorité : Arbitrage CTO", "PFU 31,4% sur les plus-values → calculer le seuil de rentabilité"),
        ("Méthode Swedroe (recommandée)", "Rebalancer si dérive > 25% de la cible (relatif) OU > 5 pts absolus"),
        ("Fréquence recommandée", "Vérification annuelle + après chaque mouvement de marché > 20%"),
        ("Coût fiscal arbitrage CTO", "PV × 31,4% → seuil : dérive > 3-5 ans de surperformance de l'allocation cible"),
    ]
    for methode, explication in methodes:
        ws.cell(row=row, column=1, value=methode).font = _font(bold=True, color=COULEUR_SUBHEADER)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row=row, column=2, value=explication).alignment = _align("left", "center", wrap=True)
        ws.row_dimensions[row].height = 22
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = _thin_border()
        row += 1

    widths = [28, 14, 14, 14, 12, 14, 22, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A5"
