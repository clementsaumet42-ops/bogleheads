import yaml
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

ROOT = Path(__file__).parent.parent.parent

# ─── Palette couleurs ───────────────────────────────────────────────
COULEURS_CLASSES = {
    "Actions": "4472C4",
    "Obligations": "ED7D31",
    "Immobilier": "A9D18E",
    "Or": "FFD966",
    "Matières premières": "9DC3E6",
    "Monétaire": "70AD47",
    "Thématiques": "BF8FBF",
    "Diversifiants": "FF9999",
}

COULEUR_HEADER = "1F3864"
COULEUR_SUBHEADER = "2E75B6"
COULEUR_AVERTISSEMENT = "C00000"
COULEUR_OK = "70AD47"
COULEUR_WARNING = "FFC000"
COULEUR_DANGER = "FF0000"
COULEUR_LIGHT_BLUE = "DDEEFF"
COULEUR_LIGHT_GREY = "F2F2F2"
COULEUR_GOLD = "FFD966"

# Aliases for shim
VERT_FONCE = "375623"
VERT_CLAIR = "C6EFCE"
BLEU_NUIT = COULEUR_HEADER
ORANGE_DOUX = "ED7D31"
GRIS_CLAIR = COULEUR_LIGHT_GREY
GRIS_MOYEN = "D9D9D9"
ROUGE_ALERTE = COULEUR_AVERTISSEMENT
BLANC = "FFFFFF"
JAUNE_OR = COULEUR_GOLD
BLEU_CIEL = COULEUR_LIGHT_BLUE


# ─── Helpers de style ───────────────────────────────────────────────
def load_yaml(filename: str) -> dict:
    with open(ROOT / "config" / filename, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(bold=bold, color=color, size=size, italic=italic)


def _align(h="left", v="center", wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _thin_border() -> Border:
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def style_header(cell, bg=COULEUR_HEADER, fg="FFFFFF", size=11, bold=True):
    cell.fill = _fill(bg)
    cell.font = Font(bold=bold, color=fg, size=size)
    cell.alignment = _align("center", "center", wrap=True)
    cell.border = _thin_border()


def style_subheader(cell, bg=COULEUR_SUBHEADER):
    cell.fill = _fill(bg)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.alignment = _align("center", "center")
    cell.border = _thin_border()


def style_data(cell, bg=None, bold=False, align_h="left", number_format=None):
    if bg:
        cell.fill = _fill(bg)
    cell.font = _font(bold=bold)
    cell.alignment = _align(align_h)
    cell.border = _thin_border()
    if number_format:
        cell.number_format = number_format


def set_col_width(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width


def ajouter_disclaimer(ws, row: int, col_start=1, col_end=10) -> int:
    disclaimer = (
        "⚠️ AVERTISSEMENT : Ce fichier est un outil pédagogique d'aide à la décision. "
        "Il ne constitue PAS un conseil en investissement au sens de la Directive MIF II. "
        "Tout conseil doit être personnalisé par un CIF/CGP agréé AMF. "
        "Paramètres fiscaux indicatifs — valider avec votre expert-comptable."
    )
    ws.merge_cells(
        start_row=row, start_column=col_start, end_row=row, end_column=col_end
    )
    cell = ws.cell(row=row, column=col_start, value=disclaimer)
    cell.fill = _fill("FFF2CC")
    cell.font = Font(bold=True, color=COULEUR_AVERTISSEMENT, size=9, italic=True)
    cell.alignment = _align("left", "center", wrap=True)
    ws.row_dimensions[row].height = 35
    return row + 1


def titre_section(ws, row: int, texte: str, col_start=1, col_end=8, bg=COULEUR_SUBHEADER) -> int:
    ws.merge_cells(
        start_row=row, start_column=col_start, end_row=row, end_column=col_end
    )
    cell = ws.cell(row=row, column=col_start, value=texte)
    cell.fill = _fill(bg)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.alignment = _align("left", "center")
    ws.row_dimensions[row].height = 20
    return row + 1
