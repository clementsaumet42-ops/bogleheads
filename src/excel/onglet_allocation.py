import openpyxl
from openpyxl.styles import Font
from src.excel.styles import (
    COULEURS_CLASSES, COULEUR_HEADER, COULEUR_SUBHEADER, COULEUR_LIGHT_GREY,
    _fill, _font, _align, _thin_border,
    style_header, style_subheader, style_data,
    set_col_width, ajouter_disclaimer, titre_section,
)


# ─── Onglet 5 : Allocation Cible ────────────────────────────────────
def creer_onglet_allocation_cible(wb: openpyxl.Workbook, profil: dict = None):
    ws = wb.create_sheet("Allocation_Cible")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "🎯 ALLOCATION CIBLE — PORTEFEUILLE BOGLEHEAD FR"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30
    ajouter_disclaimer(ws, 2, 1, 7)

    row = 4
    row = titre_section(ws, row, "📊 ALLOCATION CIBLE PAR CLASSE D'ACTIFS", 1, 7)
    headers = ["Classe d'actifs", "Allocation cible (%)", "Montant (€)", "ETF de référence", "Enveloppe prioritaire", "Justification", "Couleur"]
    for i, h in enumerate(headers, 1):
        style_subheader(ws.cell(row=row, column=i, value=h))
    row += 1

    if profil:
        alloc = profil.get("allocation_cible_bogleheads", {})
        patrimoine = profil.get("patrimoine_financier_total", 0)
        commentaire = alloc.get("commentaire", "")
        classes = {
            "actions": ("Actions", "CW8 (PEA) / IWDA (CTO)", "PEA → CTO → PER"),
            "obligations": ("Obligations", "GOVS / AGGH", "PER → Contrat Cap IS"),
            "immobilier_cote": ("Immobilier coté (REITs)", "IWDP / EPRE", "PER → CTO"),
            "or": ("Or physique (ETC)", "GOLD / IGLN", "CTO → PER"),
            "liquidites": ("Liquidités / Monétaire", "CSH / XEON", "CTO IS → CTO"),
        }
        classe_colors = {
            "actions": "4472C4",
            "obligations": "ED7D31",
            "immobilier_cote": "A9D18E",
            "or": "FFD966",
            "liquidites": "70AD47",
        }
        for key, (label, etf_ref, env_prio) in classes.items():
            pct = alloc.get(key, 0.0)
            montant = pct * patrimoine
            bg = classe_colors.get(key, "FFFFFF")
            row_data = [label, pct, montant, etf_ref, env_prio, "", ""]
            for j, val in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=j, value=val)
                if j == 1:
                    cell.fill = _fill(bg)
                    cell.font = _font(bold=True, size=10)
                elif j == 2:
                    cell.number_format = "0.0%"
                    cell.font = _font(bold=True, size=11)
                    cell.alignment = _align("center")
                elif j == 3:
                    cell.number_format = "#,##0 €"
                elif j == 7:
                    cell.fill = _fill(bg)
                cell.border = _thin_border()
            row += 1

        # Total
        ws.cell(row=row, column=1, value="TOTAL").font = _font(bold=True, size=11)
        ws.cell(row=row, column=1).fill = _fill(COULEUR_HEADER)
        ws.cell(row=row, column=1).font = Font(bold=True, color="FFFFFF", size=11)
        total_pct = sum(v for k, v in alloc.items() if k != "commentaire")
        ws.cell(row=row, column=2, value=total_pct).number_format = "0.0%"
        ws.cell(row=row, column=2).font = _font(bold=True, size=11)
        ws.cell(row=row, column=3, value=patrimoine).number_format = "#,##0 €"
        ws.cell(row=row, column=3).font = _font(bold=True)
        row += 2

        if commentaire:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            ws.cell(row=row, column=1, value=f"💬 {commentaire}")
            ws.cell(row=row, column=1).font = _font(italic=True, size=10)
            ws.cell(row=row, column=1).fill = _fill("FFF2CC")
            ws.row_dimensions[row].height = 30
            row += 1
    else:
        ws.cell(row=row, column=1, value="ℹ️ Sélectionner un profil client pour afficher l'allocation.")

    # Règles Boglehead
    row += 1
    row = titre_section(ws, row, "📏 RÈGLES BOGLEHEAD FR (HEURISTIQUES)", 1, 7)
    regles = [
        "Règle de l'âge : % obligations ≈ âge − 10 (heuristique — à adapter au profil risque)",
        "Règle 3 fonds : 1 ETF Monde + 1 ETF Obligations + 1 ETF Émergents",
        "Règle glide path : réduire progressivement les actions à l'approche de la retraite",
        "Rebalancement : méthode Larry Swedroe — si dérive > 25% de la cible (relatif) ou > 5 pts absolus",
        "Asset location : actions croissance en PEA, obligations en PER, or en CTO",
        "Frugalité : privilégier les ETF à TER < 0,20% — chaque point de frais = impact majeur sur 20 ans",
    ]
    for regle in regles:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row=row, column=1, value=f"• {regle}")
        ws.cell(row=row, column=1).font = _font(size=10)
        ws.cell(row=row, column=1).fill = _fill(COULEUR_LIGHT_GREY if row % 2 == 0 else "FFFFFF")
        ws.row_dimensions[row].height = 20
        row += 1

    widths = [35, 18, 18, 30, 28, 35, 10]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A4"
