"""
Onglet Excel Allocation_Optimisee — Sprint S2.

Affiche :
  - Mode A : allocation cible optimisée (classes d'actifs, poids, rendement/volatilité/Sharpe)
  - Mode B : ventilation par enveloppe (montants par classe × enveloppe)
  - Coût annuel optimisé vs scénario naïf (économie démonstrative)
"""

from __future__ import annotations

from typing import Any

import openpyxl
from openpyxl.styles import Font

from src.excel.styles import (
    COULEUR_HEADER,
    COULEUR_LIGHT_BLUE,
    COULEUR_LIGHT_GREY,
    COULEUR_SUBHEADER,
    _align,
    _fill,
    _font,
    _thin_border,
    ajouter_disclaimer,
    set_col_width,
    style_subheader,
    titre_section,
)
from src.optimiseur_allocation import (
    CLASSES_ACTIFS_ORDRE,
    charger_config_optimiseur,
    optimiser_portefeuille_complet,
)

VERT_OK = "C6EFCE"
ORANGE_WARN = "FFEB9C"
NB_COLS = 10

# Labels affichage par classe d'actifs
LABELS_CLASSES = {
    "actions_usa": "Actions USA",
    "actions_dev_ex_usa": "Actions Dev. ex-USA",
    "actions_em": "Actions Émergents",
    "obligations_agg_monde": "Obligations Agrégées Monde",
    "obligations_euro": "Obligations Euro",
    "reit": "Immobilier coté (REIT)",
    "or_matieres": "Or / Matières premières",
    "monetaire": "Monétaire / Liquidités",
}

COULEURS_CLASSES = {
    "actions_usa": "4472C4",
    "actions_dev_ex_usa": "5B9BD5",
    "actions_em": "70AD47",
    "obligations_agg_monde": "ED7D31",
    "obligations_euro": "FFC000",
    "reit": "A9D18E",
    "or_matieres": "FFD966",
    "monetaire": "BFBFBF",
}


# ─── Fonctions helpers ────────────────────────────────────────────────────────


def _cell(ws, row: int, col: int, value, bold=False, bg=None, size=10, fmt=None, align_h="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _font(bold=bold, size=size)
    c.alignment = _align(align_h, "center")
    c.border = _thin_border()
    if bg:
        c.fill = _fill(bg)
    if fmt:
        c.number_format = fmt
    return c


def _merge_cell(ws, row: int, col_start: int, col_end: int, value, bold=False, bg=None, size=10):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    c = ws.cell(row=row, column=col_start, value=value)
    c.font = _font(bold=bold, size=size)
    c.alignment = _align("left", "center", wrap=True)
    if bg:
        c.fill = _fill(bg)
    return c


# ─── Création de l'onglet ────────────────────────────────────────────────────


def creer_onglet_allocation_optimisee(
    wb: openpyxl.Workbook,
    profil: dict | None = None,
    config: dict | None = None,
) -> openpyxl.worksheet.worksheet.Worksheet:
    """
    Crée l'onglet Allocation_Optimisee dans le classeur wb.

    Parameters
    ----------
    wb : openpyxl.Workbook
    profil : dict
        Dict profil client (depuis profils_clients.yaml). Si None → profil de démonstration.
    config : dict, optional
        Contenu de config/optimiseur.yaml. Chargé automatiquement si None.
    """
    if config is None:
        config = charger_config_optimiseur()

    ws = wb.create_sheet("Allocation_Optimisee")
    ws.sheet_view.showGridLines = False

    # Largeurs de colonnes
    for col, width in enumerate([3, 28, 14, 14, 14, 14, 14, 14, 14, 14], start=1):
        set_col_width(ws, col, width)

    # Calcul des résultats
    if profil is None:
        profil = _profil_demonstration()

    nom_profil = profil.get("nom", "Profil inconnu")
    code_profil = profil.get("code", "")
    resultat = optimiser_portefeuille_complet(profil, config)
    res_a = resultat["resultat_mode_a"]
    res_b = resultat["resultat_mode_b"]
    poids = res_a["poids"]
    patrimoine = float(profil.get("patrimoine_financier_total", 100_000) or 100_000)

    # ── Titre principal ───────────────────────────────────────────────────────
    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NB_COLS)
    c = ws.cell(row=row, column=1, value=f"⭐ ALLOCATION OPTIMISÉE — PROFIL {nom_profil}")
    c.font = Font(bold=True, color="FFFFFF", size=14)
    c.fill = _fill(COULEUR_HEADER)
    c.alignment = _align("center", "center")
    ws.row_dimensions[row].height = 30
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NB_COLS)
    c = ws.cell(
        row=row,
        column=1,
        value=(
            f"Code : {code_profil}  |  "
            f"Profil d'aversion : {profil.get('profil_aversion_risque', 'équilibré')}  |  "
            f"Statut optimiseur : {res_a['statut']}"
        ),
    )
    c.font = _font(italic=True, size=10, color="444444")
    c.alignment = _align("center", "center")
    row += 1
    row = ajouter_disclaimer(ws, row, col_start=1, col_end=NB_COLS)
    row += 1

    # ── Mode A : Allocation cible optimisée ───────────────────────────────────
    row = titre_section(ws, row, "📊 MODE A — ALLOCATION CIBLE OPTIMISÉE (Markowitz)", 1, NB_COLS)

    # En-têtes tableau Mode A
    headers_a = ["Classe d'actifs", "Poids (%)", "Montant (€)", "TER moyen", ""]
    for i, h in enumerate(headers_a, start=2):
        c = ws.cell(row=row, column=i, value=h)
        style_subheader(c)
    row += 1

    classes_affichees = [c for c in CLASSES_ACTIFS_ORDRE if poids.get(c, 0.0) > 0.001]
    for classe in classes_affichees:
        pct = poids.get(classe, 0.0)
        montant = pct * patrimoine
        ter = config["classes_actifs"].get(classe, {}).get("frais_ter_moyen", 0.0)
        bg = COULEURS_CLASSES.get(classe, "FFFFFF")
        label = LABELS_CLASSES.get(classe, classe)

        _cell(ws, row, 2, label, bold=True, bg=bg)
        _cell(ws, row, 3, pct, fmt="0.0%", align_h="center", bg=COULEUR_LIGHT_GREY)
        _cell(ws, row, 4, montant, fmt="#,##0 €", bg=COULEUR_LIGHT_GREY)
        _cell(ws, row, 5, ter, fmt="0.000%", align_h="center", bg=COULEUR_LIGHT_GREY)
        row += 1

    # Total allocation
    ws.cell(row=row, column=2, value="TOTAL").font = Font(bold=True, color="FFFFFF", size=10)
    ws.cell(row=row, column=2).fill = _fill(COULEUR_HEADER)
    ws.cell(row=row, column=2).border = _thin_border()
    total_pct = sum(poids.get(c, 0.0) for c in classes_affichees)
    _cell(ws, row, 3, total_pct, fmt="0.0%", bold=True, align_h="center", bg=COULEUR_HEADER)
    ws.cell(row=row, column=3).font = Font(bold=True, color="FFFFFF", size=10)
    _cell(ws, row, 4, patrimoine, fmt="#,##0 €", bold=True, bg=COULEUR_HEADER)
    ws.cell(row=row, column=4).font = Font(bold=True, color="FFFFFF", size=10)
    row += 2

    # Indicateurs de performance
    row = titre_section(ws, row, "📈 INDICATEURS DE PERFORMANCE ATTENDUE", 1, NB_COLS)
    indicateurs = [
        ("Rendement attendu annuel", f"{res_a['rendement_attendu']:.1%}"),
        ("Volatilité attendue annuelle", f"{res_a['volatilite_attendue']:.1%}"),
        (
            "Ratio Sharpe (rf=2,5%)",
            f"{res_a['ratio_sharpe']:.2f}",
        ),
    ]
    for label, val in indicateurs:
        _cell(ws, row, 2, label, bold=True, bg=COULEUR_LIGHT_GREY)
        _cell(ws, row, 3, val, align_h="center", bg=COULEUR_LIGHT_BLUE)
        row += 1
    row += 1

    # ── Mode B : Ventilation par enveloppe ────────────────────────────────────
    row = titre_section(ws, row, "🏦 MODE B — VENTILATION PAR ENVELOPPE (MILP)", 1, NB_COLS)

    ventilation = res_b.get("ventilation", [])
    if ventilation:
        # Construire tableau croisé classe × enveloppe
        envs = sorted({v["enveloppe"] for v in ventilation})
        nb_envs = len(envs)

        # En-tête
        _cell(ws, row, 2, "Classe d'actifs", bold=True, bg=COULEUR_SUBHEADER)
        ws.cell(row=row, column=2).font = Font(bold=True, color="FFFFFF", size=10)
        for j, e in enumerate(envs, start=3):
            _cell(ws, row, j, e, bold=True, bg=COULEUR_SUBHEADER, align_h="center")
            ws.cell(row=row, column=j).font = Font(bold=True, color="FFFFFF", size=10)
        _cell(ws, row, 3 + nb_envs, "TOTAL", bold=True, bg=COULEUR_SUBHEADER, align_h="center")
        ws.cell(row=row, column=3 + nb_envs).font = Font(bold=True, color="FFFFFF", size=10)
        row += 1

        totaux_envs: dict[str, float] = {e: 0.0 for e in envs}
        for classe in classes_affichees:
            bg = COULEURS_CLASSES.get(classe, "FFFFFF")
            label = LABELS_CLASSES.get(classe, classe)
            _cell(ws, row, 2, label, bold=True, bg=bg)
            total_ligne = 0.0
            for j, e in enumerate(envs, start=3):
                montant_ve = sum(
                    v["montant"]
                    for v in ventilation
                    if v["classe"] == classe and v["enveloppe"] == e
                )
                _cell(
                    ws, row, j, montant_ve if montant_ve > 0 else "", fmt="#,##0 €", align_h="right"
                )
                totaux_envs[e] += montant_ve
                total_ligne += montant_ve
            _cell(
                ws,
                row,
                3 + nb_envs,
                total_ligne,
                fmt="#,##0 €",
                bold=True,
                align_h="right",
                bg=COULEUR_LIGHT_GREY,
            )
            row += 1

        # Ligne totaux
        ws.cell(row=row, column=2, value="TOTAL").font = Font(bold=True, color="FFFFFF", size=10)
        ws.cell(row=row, column=2).fill = _fill(COULEUR_HEADER)
        ws.cell(row=row, column=2).border = _thin_border()
        grand_total = 0.0
        for j, e in enumerate(envs, start=3):
            _cell(
                ws,
                row,
                j,
                totaux_envs[e],
                fmt="#,##0 €",
                bold=True,
                bg=COULEUR_HEADER,
                align_h="right",
            )
            ws.cell(row=row, column=j).font = Font(bold=True, color="FFFFFF", size=10)
            grand_total += totaux_envs[e]
        _cell(
            ws,
            row,
            3 + nb_envs,
            grand_total,
            fmt="#,##0 €",
            bold=True,
            bg=COULEUR_HEADER,
            align_h="right",
        )
        ws.cell(row=row, column=3 + nb_envs).font = Font(bold=True, color="FFFFFF", size=10)
        row += 2
    else:
        _merge_cell(
            ws, row, 2, NB_COLS, "ℹ️ Aucune ventilation calculée (Mode B)", bg=COULEUR_LIGHT_GREY
        )
        row += 2

    # ── Coûts annuels ─────────────────────────────────────────────────────────
    row = titre_section(ws, row, "💰 COÛT ANNUEL — COMPARAISON OPTIMISÉ vs NAÏF", 1, NB_COLS)

    cout_opt = res_b.get("cout_annuel_optimise", 0.0)
    cout_naif = res_b.get("cout_annuel_naif", 0.0)
    economie = res_b.get("economie_annuelle", 0.0)
    pct_patrimoine_opt = cout_opt / patrimoine if patrimoine > 0 else 0.0
    pct_patrimoine_naif = cout_naif / patrimoine if patrimoine > 0 else 0.0

    lignes_cout = [
        ("Frais TER + gestion — Scénario optimisé", cout_opt, pct_patrimoine_opt, VERT_OK),
        (
            "Frais TER + gestion — Scénario naïf (tout CTO)",
            cout_naif,
            pct_patrimoine_naif,
            ORANGE_WARN,
        ),
        (
            "⭐ Économie annuelle estimée",
            economie,
            max(0.0, pct_patrimoine_naif - pct_patrimoine_opt),
            "E2EFDA",
        ),
    ]

    headers_c = ["", "Montant annuel (€)", "% du patrimoine", ""]
    for i, h in enumerate(headers_c, start=2):
        c = ws.cell(row=row, column=i, value=h)
        style_subheader(c)
    row += 1

    for label, montant, pct, bg in lignes_cout:
        is_economie = "Économie" in label
        _cell(ws, row, 2, label, bold=is_economie, bg=bg)
        _cell(ws, row, 3, montant, fmt="#,##0.00 €", bold=is_economie, align_h="right", bg=bg)
        _cell(ws, row, 4, pct, fmt="0.000%", bold=is_economie, align_h="center", bg=bg)
        row += 1

    row += 1
    row = ajouter_disclaimer(ws, row, col_start=1, col_end=NB_COLS)
    ws.freeze_panes = "B5"

    return ws


# ─── Profil de démonstration ──────────────────────────────────────────────────


def _profil_demonstration() -> dict[str, Any]:
    """Profil de démonstration quand aucun profil n'est fourni."""
    return {
        "id": 0,
        "code": "DEMO",
        "nom": "Profil Démonstration",
        "age": 45,
        "tmi": 0.41,
        "patrimoine_financier_total": 500_000,
        "profil_aversion_risque": "dynamique",
        "contraintes_personnalisees": {
            "exposition_usa_max": 0.50,
            "exposition_em_max": 0.15,
        },
        "enveloppes_disponibles": {
            "PEA": {"encours_actuel": 80_000, "plafond": 150_000, "ouvert": True},
            "PER": {"encours_actuel": 50_000, "ouvert": True},
            "CTO": {"encours_actuel": 150_000, "ouvert": True},
        },
    }
