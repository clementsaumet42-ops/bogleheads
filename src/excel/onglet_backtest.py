"""Onglet Excel Backtest_Comparatif — Sprint S9."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_PRIMARY = "1a4d8f"
_ACCENT = "d4a017"
_LIGHT = "f0f4fa"


def creer_onglet_backtest(wb: openpyxl.Workbook) -> None:
    """Crée l'onglet Backtest_Comparatif dans le classeur Excel."""
    ws = wb.create_sheet("Backtest_Comparatif")

    # En-tête principal
    ws.merge_cells("A1:H1")
    cell = ws["A1"]
    cell.value = "Backtest Comparatif — Portefeuilles Bogle (2003–2024)"
    cell.font = Font(bold=True, color="FFFFFF", size=13)
    cell.fill = PatternFill("solid", fgColor=_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Sous-titre
    ws.merge_cells("A2:H2")
    ws["A2"].value = (
        "Comparaison brut / net frais / net fiscal CTO / net optimisé — Données synthétiques 2003–2024"
    )
    ws["A2"].font = Font(italic=True, color="666666", size=10)
    ws["A2"].alignment = Alignment(horizontal="center")

    # En-têtes colonnes
    headers = [
        "Portefeuille",
        "Mode",
        "Capital Initial (€)",
        "Capital Final (€)",
        "CAGR (%)",
        "Volatilité Ann. (%)",
        "Sharpe",
        "Max Drawdown (%)",
    ]
    header_fill = PatternFill("solid", fgColor=_ACCENT)
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[4].height = 30

    # Données exemples (portefeuilles × modes)
    portefeuilles = [
        "BOGLE_2_FUNDS_70_30",
        "BOGLE_3_FUNDS_60_30_10",
        "BOGLE_4_FUNDS",
        "LAZY_PERMANENT_PORTFOLIO",
        "ALL_WEATHER_RAY_DALIO",
    ]
    modes = ["brut", "net_frais", "net_fiscal_cto", "net_optimise"]

    row = 5
    for i, pf in enumerate(portefeuilles):
        for j, mode in enumerate(modes):
            fill = PatternFill("solid", fgColor=_LIGHT) if (i + j) % 2 == 0 else None
            ws.cell(row=row, column=1, value=pf).alignment = Alignment(horizontal="left")
            ws.cell(row=row, column=2, value=mode)
            ws.cell(row=row, column=3, value=100_000).number_format = "#,##0.00"
            ws.cell(row=row, column=4, value="—")
            ws.cell(row=row, column=5, value="—")
            ws.cell(row=row, column=6, value="—")
            ws.cell(row=row, column=7, value="—")
            ws.cell(row=row, column=8, value="—")
            if fill:
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = fill
            row += 1

    # Note explicative
    row += 1
    ws.merge_cells(f"A{row}:H{row}")
    ws[f"A{row}"].value = (
        "ℹ️  Pour obtenir les vraies métriques, lancer build_backtest.py "
        "puis recharger ce fichier Excel."
    )
    ws[f"A{row}"].font = Font(italic=True, color="888888", size=9)

    # Largeurs colonnes
    col_widths = [30, 18, 18, 18, 12, 18, 12, 18]
    for col, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
