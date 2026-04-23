import datetime

import openpyxl
from openpyxl.styles import Font

from src.excel.styles import (
    COULEUR_HEADER,
    COULEUR_LIGHT_GREY,
    _fill,
    _font,
    _thin_border,
    ajouter_disclaimer,
    set_col_width,
    style_header,
    style_subheader,
    titre_section,
)


# ─── Onglet 8 : Reporting Client ────────────────────────────────────
def creer_onglet_reporting(wb: openpyxl.Workbook, profil: dict = None):
    ws = wb.create_sheet("Reporting_Client")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "📋 REPORTING CLIENT — BOGLEHEAD FR"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30

    # Date de génération
    ws.cell(row=2, column=1, value=f"Généré le : {datetime.date.today().strftime('%d/%m/%Y')}")
    ws.cell(row=2, column=1).font = _font(italic=True, size=9)
    row = ajouter_disclaimer(ws, 3, 1, 8)

    if profil:
        row = titre_section(ws, row, f"👤 SYNTHÈSE — {profil.get('nom', 'Client')}", 1, 8)
        ws.cell(row=row, column=1, value="Patrimoine financier total").font = _font(bold=True)
        ws.cell(
            row=row, column=2, value=profil.get("patrimoine_financier_total", 0)
        ).number_format = "#,##0 €"
        ws.cell(row=row, column=3, value="Capacité d'épargne annuelle").font = _font(bold=True)
        ws.cell(
            row=row, column=4, value=profil.get("capacite_epargne_annuelle", 0)
        ).number_format = "#,##0 €"
        row += 1
        ws.cell(row=row, column=1, value="Horizon de placement").font = _font(bold=True)
        ws.cell(row=row, column=2, value=f"{profil.get('horizon_placement_ans', 0)} ans")
        ws.cell(row=row, column=3, value="Score de risque SRRI").font = _font(bold=True)
        ws.cell(row=row, column=4, value=f"{profil.get('score_risque', '—')} / 7")
        row += 2

        row = titre_section(ws, row, "🏦 RÉPARTITION PAR ENVELOPPE", 1, 8)
        headers = [
            "Enveloppe",
            "Encours actuel (€)",
            "Versements prévus/an (€)",
            "Plafond restant (€)",
            "Avantage fiscal clé",
            "",
            "",
            "",
        ]
        for i, h in enumerate(headers, 1):
            style_subheader(ws.cell(row=row, column=i, value=h))
        row += 1

        env_dispo = profil.get("enveloppes_disponibles", {})
        env_labels = {
            "PEA": ("PEA", "Exonération IR après 5 ans"),
            "PER": ("PER (Plan Épargne Retraite)", "Déduction fiscale à l'entrée"),
            "PEE": ("PEE (Plan Épargne Entreprise)", "Abondement + exonération IR"),
            "CTO_perso": ("CTO Personnel", "Liquidité totale"),
            "CTO_IS": ("CTO IS (Holding)", "Taux IS 15/25%"),
            "Contrat_Cap_IS": ("Contrat Cap. IS (Holding)", "Pas de mark-to-market"),
        }
        total_patrimoine_env = 0
        for env_key, (label, avantage) in env_labels.items():
            env_data = env_dispo.get(env_key)
            if env_data is None:
                continue
            encours = env_data.get("encours_actuel", 0) or 0
            versements = env_data.get("versement_annuel_prevu", 0) or 0
            plafond = env_data.get("plafond", None)
            plafond_restant = max(0, plafond - encours) if plafond else "Illimité"
            total_patrimoine_env += encours

            row_data = [label, encours, versements, plafond_restant, avantage, "", "", ""]
            for j, val in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=j, value=val)
                if (
                    j == 2
                    and isinstance(val, (int, float))
                    or j == 3
                    and isinstance(val, (int, float))
                    or j == 4
                    and isinstance(val, (int, float))
                ):
                    cell.number_format = "#,##0 €"
                cell.border = _thin_border()
            row += 1

        ws.cell(row=row, column=1, value="TOTAL enveloppes").font = _font(bold=True)
        ws.cell(row=row, column=1).fill = _fill(COULEUR_HEADER)
        ws.cell(row=row, column=1).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=row, column=2, value=total_patrimoine_env).number_format = "#,##0 €"
        ws.cell(row=row, column=2).font = _font(bold=True)
        row += 2

        row = titre_section(ws, row, "🎯 OBJECTIFS ET PRÉCONISATIONS", 1, 8)
        ws.cell(row=row, column=1, value="Objectif principal").font = _font(bold=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row=row, column=2, value=profil.get("objectif_principal", "—"))
        ws.row_dimensions[row].height = 20
        row += 1
        for i, obj in enumerate(profil.get("objectifs_secondaires", []), 1):
            ws.cell(row=row, column=1, value=f"Obj. secondaire {i}").font = _font(bold=True)
            ws.cell(row=row, column=1).fill = _fill(COULEUR_LIGHT_GREY)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            ws.cell(row=row, column=2, value=obj)
            row += 1
    else:
        ws.cell(row=row, column=1, value="ℹ️ Sélectionner un profil pour générer le reporting.")

    widths = [35, 20, 22, 18, 35, 15, 15, 15]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A4"
