import openpyxl
from openpyxl.styles import Font
from src.excel.styles import (
    COULEURS_CLASSES, COULEUR_SUBHEADER, COULEUR_LIGHT_GREY,
    _fill, _font, _align, _thin_border,
    style_header, style_subheader, set_col_width, ajouter_disclaimer, titre_section,
)


# ─── Onglet 6 : Asset Location Matrice ──────────────────────────────
def creer_onglet_asset_location(wb: openpyxl.Workbook, etfs: list, enveloppes: list, profil: dict = None):
    ws = wb.create_sheet("Asset_Location_Matrice")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = "🗺️ MATRICE D'ASSET LOCATION — ETF × ENVELOPPES"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30
    ajouter_disclaimer(ws, 2, 1, 9)

    row = 4
    row = titre_section(ws, row, "📊 MATRICE ÉLIGIBILITÉ ETF × ENVELOPPE", 1, 9)

    env_ids = ["PEA", "PER", "PEE", "CTO_perso", "CTO_IS", "Contrat_Cap_IS"]
    headers = ["ISIN", "Nom ETF", "Classe"] + env_ids
    for i, h in enumerate(headers, 1):
        style_subheader(ws.cell(row=row, column=i, value=h))
    row += 1

    for etf in etfs:
        elig = etf.get("eligibilite", {})
        classe = etf.get("classe_actifs", "")
        bg_hex = COULEURS_CLASSES.get(classe, "FFFFFF")
        row_data = [etf.get("isin", ""), etf.get("nom", ""), classe]
        for env_id in env_ids:
            row_data.append("✅" if elig.get(env_id) else "—")

        for j, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=j, value=val)
            cell.font = _font(size=9)
            cell.border = _thin_border()
            if j <= 3:
                cell.fill = _fill(bg_hex)
                cell.alignment = _align("left", "center", wrap=True)
            else:
                cell.alignment = _align("center")
                if val == "✅":
                    cell.fill = _fill("C6EFCE")
                else:
                    cell.fill = _fill("FCE4D6")
        row += 1

    row += 1
    # Légende
    row = titre_section(ws, row, "🔑 RÈGLES D'ASSET LOCATION BOGLEHEAD FR", 1, 9)
    regles_al = [
        ("Actions mondiales (PEA éligible)", "PEA en PRIORITÉ → CTO → PER", "Économie IR 12,8% sur la totalité des gains"),
        ("Actions hors PEA (ETF physiques)", "PER → CTO perso → Contrat Cap IS", "PER = déduction entrée, CTO = liquidité"),
        ("Obligations / Fixed Income", "PER → Contrat Cap IS → CTO IS", "Rendement fixe mieux dans enveloppe défiscalisée"),
        ("Or (ETCs)", "CTO perso → PER", "ETCs non éligibles PEA — CTO simple"),
        ("REITs / Immobilier coté", "PER → CTO perso", "Dividendes imposables — mieux dans enveloppe"),
        ("Monétaire / Liquidités", "CTO IS → CTO perso", "Faible rendement — hors enveloppes précieuses"),
        ("ETF actions IS (holding)", "Contrat Cap IS > CTO IS", "Éviter le mark-to-market annuel (art. 209-0 A CGI)"),
    ]
    ws.cell(row=row, column=1, value="Classe d'actifs").font = _font(bold=True)
    ws.cell(row=row, column=2, value="Enveloppe recommandée").font = _font(bold=True)
    ws.cell(row=row, column=3, value="Justification").font = _font(bold=True)
    for j in range(1, 4):
        ws.cell(row=row, column=j).fill = _fill(COULEUR_SUBHEADER)
        ws.cell(row=row, column=j).font = Font(bold=True, color="FFFFFF")
    row += 1
    for i, (classe, reco, justif) in enumerate(regles_al):
        bg = COULEUR_LIGHT_GREY if i % 2 == 0 else "FFFFFF"
        ws.cell(row=row, column=1, value=classe).fill = _fill(bg)
        ws.cell(row=row, column=2, value=reco).fill = _fill(bg)
        ws.cell(row=row, column=2).font = _font(bold=True, color=COULEUR_SUBHEADER)
        ws.cell(row=row, column=3, value=justif).fill = _fill(bg)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
        ws.cell(row=row, column=3).alignment = _align("left", "center", wrap=True)
        row += 1

    widths = [16, 45, 18, 8, 8, 8, 8, 8, 18]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A5"
