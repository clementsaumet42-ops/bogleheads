from pathlib import Path

import openpyxl

from src.excel.onglet_allocation import creer_onglet_allocation_cible
from src.excel.onglet_allocation_optimisee import creer_onglet_allocation_optimisee
from src.excel.onglet_asset_location import creer_onglet_asset_location
from src.excel.onglet_backtest import creer_onglet_backtest
from src.excel.onglet_enveloppes import creer_onglet_enveloppes
from src.excel.onglet_glide_path import _creer_onglet_glide_path
from src.excel.onglet_parametres import creer_onglet_fiscalite, creer_onglet_parametres_client
from src.excel.onglet_plan_rebalancement import creer_onglet_plan_rebalancement
from src.excel.onglet_profils import (
    creer_onglet_comparatif_profils,
    creer_onglet_profil_individuel,
    creer_onglet_profils_types,
)
from src.excel.onglet_projection import _creer_onglet_projection_monte_carlo
from src.excel.onglet_rebalancement import creer_onglet_rebalancement
from src.excel.onglet_rebalancement_flux import _creer_onglet_rebalancement_flux
from src.excel.onglet_reporting import creer_onglet_reporting
from src.excel.onglet_tuto_solveur import creer_onglet_tuto_solveur
from src.excel.onglet_univers_etf import creer_onglet_univers_etf
from src.excel.onglet_alertes import creer_onglet_alertes
from src.excel.styles import ROOT, load_yaml


def generer_excel(chemin_sortie: str = None):
    """Génère le fichier Excel complet Boglehead FR."""
    if chemin_sortie is None:
        chemin_sortie = str(ROOT / "output" / "portefeuille_bogleheads.xlsx")

    params_fiscaux = load_yaml("fiscalite_2026.yaml")
    enveloppes_data = load_yaml("enveloppes.yaml")["enveloppes"]
    etfs_data = load_yaml("univers_etf.yaml")["univers_etf"]
    profils_data = load_yaml("profils_clients.yaml")

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    print("  → Onglet Paramètres_Client")
    creer_onglet_parametres_client(wb)

    print("  → Onglet Paramètres_Fiscalité_2026")
    creer_onglet_fiscalite(wb, params_fiscaux)

    print("  → Onglet Enveloppes")
    creer_onglet_enveloppes(wb, enveloppes_data)

    print("  → Onglet Univers_ETF")
    creer_onglet_univers_etf(wb, etfs_data)

    print("  → Onglet Allocation_Cible")
    creer_onglet_allocation_cible(wb)

    print("  → Onglet Asset_Location_Matrice")
    creer_onglet_asset_location(wb, etfs_data, enveloppes_data)

    print("  → Onglet Rebalancement")
    creer_onglet_rebalancement(wb)

    print("  → Onglet Reporting_Client")
    creer_onglet_reporting(wb)

    print("  → Onglet Profils_Types")
    creer_onglet_profils_types(wb, profils_data)

    for profil in profils_data.get("profils", []):
        nom_court = profil.get("nom", f"Profil_{profil['id']}")
        print(f"  → Onglet Profil {profil['id']} — {nom_court}")
        creer_onglet_profil_individuel(wb, profil, etfs_data, params_fiscaux)

    print("  → Onglet Comparatif_Profils")
    creer_onglet_comparatif_profils(wb, profils_data, params_fiscaux)

    print("  → Onglet Tuto_Solveur")
    creer_onglet_tuto_solveur(wb)

    print("  → Onglet Projection_MonteCarlo")
    profil_ref = profils_data.get("profils", [{}])[0]
    _creer_onglet_projection_monte_carlo(wb, profil_ref)

    print("  → Onglet Glide_Path")
    _creer_onglet_glide_path(wb, profil_ref)

    print("  → Onglet Rebalancement_Flux")
    _creer_onglet_rebalancement_flux(wb, profil_ref)

    print("  → Onglet Plan_Rebalancement")
    creer_onglet_plan_rebalancement(wb, profil_ref)

    print("  → Onglet Allocation_Optimisee")
    creer_onglet_allocation_optimisee(wb, profil_ref)

    print("  → Onglet Backtest_Comparatif")
    creer_onglet_backtest(wb)

    print("  → Onglet Alertes_40_Règles")
    try:
        from src.audit.alertes import detecter_alertes
        from src.schemas import Profil as _ProfilSchema
        _profil_ref_obj = _ProfilSchema.model_validate(profil_ref)
        _alertes = detecter_alertes(_profil_ref_obj)
        _ws_alertes = wb.create_sheet("Alertes (40 règles)")
        creer_onglet_alertes(_ws_alertes, _alertes)
    except Exception as _exc:
        print(f"     ⚠ Onglet alertes ignoré : {_exc}")

    wb.properties.title = "Boglehead FR — Outil CGP Multi-Enveloppes 2026"
    wb.properties.subject = "Allocation Boglehead multi-enveloppes — France 2026"
    wb.properties.creator = "Boglehead FR — Outil CGP"
    wb.properties.description = (
        "Outil pédagogique de gestion de portefeuille Boglehead multi-enveloppes fiscales. "
        "Ne constitue pas un conseil en investissement."
    )

    Path(chemin_sortie).parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin_sortie)
    print(f"\n✅ Fichier sauvegardé : {chemin_sortie}")
    print(f"   Onglets créés : {len(wb.sheetnames)}")
    for name in wb.sheetnames:
        print(f"   • {name}")
