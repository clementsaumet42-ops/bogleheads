import openpyxl

from src.excel.styles import (
    COULEUR_SUBHEADER,
    _align,
    _fill,
    _font,
    _thin_border,
    ajouter_disclaimer,
    set_col_width,
    style_header,
    style_subheader,
    titre_section,
)


# ─── Onglet 3 : Enveloppes ──────────────────────────────────────────
def creer_onglet_enveloppes(wb: openpyxl.Workbook, enveloppes: list):
    ws = wb.create_sheet("Enveloppes")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "🏦 ENVELOPPES FISCALES DISPONIBLES — FRANCE 2026"
    style_header(c, size=14)
    ws.row_dimensions[1].height = 30

    row = ajouter_disclaimer(ws, 2, 1, 8)

    headers = [
        "ID Enveloppe",
        "Nom complet",
        "Plafond (€)",
        "Fiscalité sortie (résumé)",
        "Avantages clés",
        "Inconvénients",
        "ETF éligibles",
        "Notes CGP",
    ]
    row = titre_section(ws, row, "📋 TABLEAU RÉCAPITULATIF DES ENVELOPPES", 1, 8)
    for i, h in enumerate(headers, 1):
        style_subheader(ws.cell(row=row, column=i, value=h))
    row += 1

    env_colors = {
        "PEA": "C6EFCE",
        "PER": "BDD7EE",
        "PEE": "FFEB9C",
        "CTO_perso": "F2F2F2",
        "CTO_IS": "FCE4D6",
        "Contrat_Cap_IS": "E2EFDA",
    }

    for env in enveloppes:
        bg = env_colors.get(env["id"], "FFFFFF")
        plafond = env.get("plafond") or "Illimité"
        if isinstance(plafond, (int, float)):
            plafond = f"{plafond:,} €".replace(",", " ")

        # Sortie fiscale résumée
        sortie = env.get("fiscalite_sortie", "")
        if isinstance(sortie, dict):
            sortie = " | ".join(f"{k}: {v}" for k, v in sortie.items())

        avantages = "\n".join(f"• {a}" for a in env.get("avantages", []))
        inconvenients = "\n".join(f"• {i}" for i in env.get("inconvenients", []))
        etfs = ", ".join(env.get("eligible_etf_types", []))

        row_data = [
            env["id"],
            env["nom"],
            plafond,
            sortie,
            avantages,
            inconvenients,
            etfs,
            env.get("notes", ""),
        ]
        max_lines = max(len(avantages.split("\n")), len(inconvenients.split("\n")), 2)
        ws.row_dimensions[row].height = max(25, max_lines * 15)

        for j, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=j, value=val)
            cell.fill = _fill(bg)
            cell.font = _font(size=9, bold=(j == 1))
            cell.alignment = _align("left", "top", wrap=True)
            cell.border = _thin_border()
        row += 1

    # Priorités d'utilisation
    row += 1
    row = titre_section(ws, row, "🎯 ORDRE DE PRIORITÉ D'UTILISATION (BOGLEHEAD FR)", 1, 8)
    priorites = [
        (
            "1️⃣",
            "PEE",
            "TOUJOURS saturer l'abondement employeur en PREMIER — TRI immédiat imbattable",
        ),
        (
            "2️⃣",
            "PEA",
            "Priorité croissance long terme — exonération IR après 5 ans — plafond 150 000 €",
        ),
        ("3️⃣", "PER", "Si TMI actuelle > TMI retraite estimée — déduction fiscale à l'entrée"),
        (
            "4️⃣",
            "Contrat Cap IS",
            "Si holding IS — pour les ETF capitalisants — pas de mark-to-market",
        ),
        ("5️⃣", "CTO perso", "Surplus d'épargne — liquidité maximale — PFU 31,4%"),
        ("⚠️", "CTO IS", "PIÈGE : mark-to-market annuel — préférer contrat cap IS pour les ETF"),
    ]
    for prio, env_id, explication in priorites:
        ws.cell(row=row, column=1, value=prio).font = _font(bold=True, size=12)
        ws.cell(row=row, column=2, value=env_id).font = _font(bold=True, color=COULEUR_SUBHEADER)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row=row, column=3, value=explication).alignment = _align(
            "left", "center", wrap=True
        )
        ws.row_dimensions[row].height = 20
        row += 1

    widths = [18, 35, 14, 40, 40, 40, 35, 40]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.freeze_panes = "A4"
